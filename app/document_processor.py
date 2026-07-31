"""
app/document_processor.py — Učitavanje, analiza i segmentacija dokumenata
Ref: doc/README_TechDoc.md §7
"""

from __future__ import annotations

import os
import re
import logging
from collections import Counter
from typing import Any

from pypdf import PdfReader
from docx import Document

# Pokušaj uvoza za EPUB / MOBI - neobavezni paketi
try:
    import ebooklib
    from ebooklib import epub
    _EPUB_AVAILABLE = True
except ImportError:
    _EPUB_AVAILABLE = False

try:
    import mobi
    _MOBI_AVAILABLE = True
except ImportError:
    _MOBI_AVAILABLE = False


class DocumentProcessor:
    """Klasa za učitavanje, analizu i segmentaciju dokumenata."""

    def __init__(self, config: dict[str, Any]) -> None:
        self._cfg = config
        self._scan_limit = config.get("translation", {}).get("scan_pages_limit", 30)
        self._hf_threshold = config.get("translation", {}).get("header_footer_threshold", 0.40)
        self._chapter_patterns = config.get("chapter_patterns", [])

    def ucitaj_izvorni_tekst(self, putanja_knjige: str) -> str:
        """Učitava sirovi tekst iz datoteke bez frekvencijske analize (za već očišćene knjige).

        Args:
            putanja_knjige: Apsolutna putanja do datoteke.

        Returns:
            Sirovi tekst iz datoteke.
        """
        ekstenzija = os.path.splitext(putanja_knjige)[1].lower()
        logging.info(f"Učitavam izvorni tekst: {os.path.basename(putanja_knjige)}")

        if ekstenzija == '.txt':
            with open(putanja_knjige, 'r', encoding='utf-8') as f:
                return f.read()
        elif ekstenzija == '.docx':
            doc = Document(putanja_knjige)
            return "\n".join([p.text for p in doc.paragraphs])
        elif ekstenzija == '.pdf':
            reader = PdfReader(putanja_knjige)
            dijelovi: list[str] = []
            for stranica in reader.pages:
                t = stranica.extract_text()
                if t:
                    dijelovi.append(t)
            return "\n".join(dijelovi)
        elif ekstenzija == '.epub':
            if not _EPUB_AVAILABLE:
                logging.warning(
                    "Biblioteka 'ebooklib' nije instalirana. "
                    "Instalirajte je s: pip install ebooklib"
                )
                raise ImportError(
                    "Biblioteka 'ebooklib' nije instalirana. "
                    "Instalirajte je s: pip install ebooklib"
                )
            knjiga = epub.read_epub(putanja_knjige)
            dijelovi = []
            for item in knjiga.get_items():
                if item.get_type() == ebooklib.ITEM_DOCUMENT:
                    content = item.get_content()
                    if content:
                        dekodirani = content.decode('utf-8', errors='replace')
                        cisti = re.sub(r'<[^>]+>', '', dekodirani)
                        dijelovi.append(cisti)
            return "\n".join(dijelovi)
        elif ekstenzija == '.mobi':
            if not _MOBI_AVAILABLE:
                logging.warning(
                    "Biblioteka 'mobi' nije instalirana. "
                    "Instalirajte je s: pip install mobi"
                )
                raise ImportError(
                    "Biblioteka 'mobi' nije instalirana. "
                    "Instalirajte je s: pip install mobi"
                )
            temp_dir, file_path = mobi.extract(putanja_knjige)
            with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
                return f.read()
        else:
            raise ValueError(f"Nepodržani format datoteke: {ekstenzija}")

    def analiziraj_i_izvuci_tekst(self, putanja_knjige: str) -> tuple[str, dict[str, list[str]], list]:
        """Provodi frekvencijsku analizu otiska te čisti Header/Footer bez hardkodiranja.

        Args:
            putanja_knjige: Apsolutna putanja do datoteke za obradu.

        Returns:
            tuple: (očišćeni_tekst, detektirani_elementi, pages)
        """
        ekstenzija = os.path.splitext(putanja_knjige)[1].lower()

        logging.info(f"Započinje analiza dokumenta: {os.path.basename(putanja_knjige)}")

        if ekstenzija == '.txt':
            with open(putanja_knjige, 'r', encoding='utf-8') as f:
                return f.read(), {}, []
        elif ekstenzija == '.docx':
            doc = Document(putanja_knjige)
            pun_tekst = "\n".join([p.text for p in doc.paragraphs])
            return pun_tekst, {}, []

        if ekstenzija in ('.epub', '.mobi'):
            try:
                pun_tekst = self.ucitaj_izvorni_tekst(putanja_knjige)
                return pun_tekst, {}, []
            except ImportError as e:
                logging.error(str(e))
                raise

        if ekstenzija != '.pdf':
            raise ValueError(f"Nepodržani format datoteke: {ekstenzija}")

        reader = PdfReader(putanja_knjige)
        ukupno_stranica = len(reader.pages)
        limit_skeniranja = min(ukupno_stranica, self._scan_limit)

        vrhovi_stranica: list[str] = []
        dna_stranica: list[str] = []

        for i in range(limit_skeniranja):
            tekst_stranice = reader.pages[i].extract_text()
            if not tekst_stranice:
                continue
            redovi = [r.strip() for r in tekst_stranice.split('\n') if r.strip()]
            if redovi:
                vrhovi_stranica.append(redovi[0])
                if len(redovi) > 1:
                    dna_stranica.append(redovi[-1])

        brojac_vrh = Counter(vrhovi_stranica)
        brojac_dno = Counter(dna_stranica)

        za_uklanjanje: list[str] = []
        detektirani_elementi: dict[str, list[str]] = {"headers": [], "footers": []}

        for tekst, count in brojac_vrh.items():
            if count / limit_skeniranja > self._hf_threshold and len(tekst) > 3:
                if not re.match(r'^\d+$', tekst):
                    za_uklanjanje.append(tekst)
                    detektirani_elementi["headers"].append(tekst)

        for tekst, count in brojac_dno.items():
            if count / limit_skeniranja > self._hf_threshold and len(tekst) > 3:
                if not re.match(r'^\d+$', tekst):
                    za_uklanjanje.append(tekst)
                    detektirani_elementi["footers"].append(tekst)

        ocisceni_tekst_lista: list[str] = []
        regex_brojevi = r'^\d+$|^\b(Page|page)\b\s*\d+'

        for i in range(ukupno_stranica):
            tekst_stranice = reader.pages[i].extract_text()
            if not tekst_stranice:
                continue

            redovi = tekst_stranice.split('\n')
            novi_redovi: list[str] = []
            for r in redovi:
                r_clean = r.strip()
                if r_clean in za_uklanjanje:
                    continue
                if re.match(regex_brojevi, r_clean):
                    continue
                novi_redovi.append(r)

            ocisceni_tekst_lista.append("\n".join(novi_redovi))

        kompletan_tekst = "\n".join(ocisceni_tekst_lista)
        return kompletan_tekst, detektirani_elementi, reader.pages

    def segmentiraj_poglavlja(self, tekst: str) -> list[dict[str, str]]:
        """Razbija očišćeni tekst na poglavlja koristeći konfiguracijske uzorke.

        Čuva strukturu odlomaka unutar svakog poglavlja (\\n prijelomi se zadržavaju).

        Args:
            tekst: Očišćeni tekst knjige.

        Returns:
            Lista poglavlja, svako s "naslov" i "sadrzaj" (s očuvanim odlomcima).
        """
        redovi = tekst.split('\n')
        poglavlja: list[dict[str, str]] = []
        trenutno_poglavlje_naziv = "Prologue/Pre-Chapter Text"
        trenutni_tekst: list[str] = []

        uzorci = [re.compile(pat) for pat in self._chapter_patterns]

        for red in redovi:
            red_clean = red.strip()
            je_naslov = False

            for uzorak in uzorci:
                if uzorak.match(red_clean):
                    je_naslov = True
                    break

            if je_naslov:
                if trenutni_tekst:
                    poglavlja.append({
                        "naslov": trenutno_poglavlje_naziv,
                        "sadrzaj": "\n".join(trenutni_tekst).strip()
                    })
                trenutno_poglavlje_naziv = red_clean
                trenutni_tekst = []
            else:
                trenutni_tekst.append(red)

        if trenutni_tekst:
            poglavlja.append({
                "naslov": trenutno_poglavlje_naziv,
                "sadrzaj": "\n".join(trenutni_tekst).strip()
            })

        if len(poglavlja) == 1 and poglavlja[0]["naslov"] == "Prologue/Pre-Chapter Text":
            logging.warning("Nije detektirana struktura poglavlja preko Regexa. Pokrećem automatsko rezanje.")
            poglavlja = []
            velicina_bloka = 15000
            pun_tekst = "\n".join(trenutni_tekst) if trenutni_tekst else tekst
            za_rezanje = [pun_tekst[i:i + velicina_bloka] for i in range(0, len(pun_tekst), velicina_bloka)]
            for idx, blok in enumerate(za_rezanje):
                poglavlja.append({
                    "naslov": f"Dio {idx + 1}",
                    "sadrzaj": blok.strip()
                })

        return poglavlja
